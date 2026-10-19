import sys, json, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from parse import parse_all, display, strip_acc, COLOUR
from desc import D; from desc_en import E
vv=json.load(open('vinvigne.json')); aocs=parse_all(); g={}; uses={}
for a in aocs:
    for w in a['wines']:
        for x in w['grapes']: g.setdefault(x['id'],x); uses.setdefault(x['id'],set()).add(a['name'])
wb=openpyxl.Workbook(); ws=wb.active; ws.title='Cépages A-Z'
ws.append(['Cépage','Couleur','Code','Nb AOC','Synonymes','Origine','Description','Description (EN)'])
for gid,x in sorted(g.items(), key=lambda kv: strip_acc(display(kv[1]['key'],kv[1]['code'])).lower()):
    i=vv.get(gid,{}); ws.append([display(x['key'],x['code']),COLOUR[x['code']],x['code'],len(uses[gid]),i.get('syn',''),i.get('origin',''),D.get(gid,''),E.get(gid,'')])
w2=wb.create_sheet('AOC × cépages')
w2.append(['AOC','Région','Type de vin','Cépage','Rôle','Proportion','Mineur (italique)','Monocépage',"Variété d'adaptation",'Source CDC'])
R={'P':'principal','C':'complémentaire','S':'secondaire','A':'accessoire','':'autorisé'}
for a in sorted(aocs,key=lambda a: strip_acc(a['name']).lower()):
    for w in a['wines']:
        for x in w['grapes']:
            w2.append([a['name'],a['region'],w['label'],display(x['key'],x['code']),R[x['role']],x['pct'],'oui' if (x['role'] in ('S','A') or x['adapt']) else '','oui' if len(w['grapes'])==1 else '','oui' if x['adapt'] else '',a['cdc']])
for s,widths in [(ws,[30,10,6,8,50,30,80,80]),(w2,[30,16,30,30,14,24,10,10,10,40])]:
    for c in s[1]: c.font=Font(bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='6B1E2E')
    for i,wd in enumerate(widths,1): s.column_dimensions[get_column_letter(i)].width=wd
    s.freeze_panes='B2'; s.auto_filter.ref=s.dimensions
for row in ws.iter_rows(min_row=2):
    for c in row[4:]: c.alignment=Alignment(wrap_text=True,vertical='top')
wb.save(sys.argv[1]); print('saved',ws.max_row-1,'grapes,',w2.max_row-1,'rows')
